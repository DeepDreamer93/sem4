#!/usr/bin/env python3
"""Build the viva deck (white / black / grey / red). Speaker notes come from Viva_Speaker_Notes.md."""
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
CH = os.path.join(HERE, 'charts')
OUT = os.path.join(HERE, 'out')
BLACK, GREY, LGREY, RED, WHITE = (RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x6E, 0x6E, 0x6E),
                                  RGBColor(0xD9, 0xD9, 0xD9), RGBColor(0xC0, 0x00, 0x00),
                                  RGBColor(0xFF, 0xFF, 0xFF))
FONT = 'Arial'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def load_notes():
    text = open(os.path.join(HERE, 'Viva_Speaker_Notes.md'), encoding='utf-8').read()
    notes = {}
    for m in re.finditer(r'^## (Slide \d+|Backup B\d) – [^\n]*\n(.*?)(?=^## |\Z)', text, re.S | re.M):
        notes[m.group(1)] = re.sub(r'\n-{3,}\s*$', '', m.group(2).strip()).strip()
    return notes


NOTES = load_notes()


def runs(par, text, size, color, bold=False):
    for tok in re.split(r'(\[\[.*?\]\]|\*\*.*?\*\*)', text):
        if not tok:
            continue
        r = par.add_run()
        c, b, t = color, bold, tok
        if tok.startswith('[['):
            c, b, t = RED, True, tok[2:-2]
        elif tok.startswith('**'):
            b, t = True, tok[2:-2]
        r.text = t
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), b
        r.font.color.rgb = c


def text(s, x, y, w, h, lines, size=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
         gap=8, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, a, Inches(0.05))
    for i, ln in enumerate([lines] if isinstance(lines, str) else lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.space_after = Pt(gap)
        sz = size
        if isinstance(ln, tuple):
            ln, sz = ln
        runs(par, ln, sz, color, bold)
    return tb


def box(s, x, y, w, h, txt, fill=None, line=GREY, size=16, color=BLACK, bold=False,
        shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.CENTER, lw=1.25):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for a in ('margin_left', 'margin_right'):
        setattr(tf, a, Inches(0.08))
    for i, ln in enumerate(txt.split('\n') if txt else ['']):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        runs(par, ln, size, color, bold)
    return sh


def img(s, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw['width'] = Inches(w)
    if h:
        kw['height'] = Inches(h)
    return s.shapes.add_picture(os.path.join(CH, name), Inches(x), Inches(y), **kw)


def table(s, x, y, w, rows, colw, size=12, row_h=0.4, red_rows=(), header=True):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                             Inches(row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    for i, cw in enumerate(colw):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LGREY if (header and r == 0) else WHITE
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            col = RED if r in red_rows else BLACK
            runs(par, val, size, col, bold=(header and r == 0))
    return tbl


COUNTER = {'n': 0}


def slide(title, key=None, number=True):
    s = prs.slides.add_slide(BLANK)
    if title:
        text(s, 0.6, 0.35, 12.1, 0.95, title, size=28, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        box(s, 0.62, 1.33, 1.3, 0.06, '', fill=RED, line=None)
    text(s, 0.6, 7.0, 8, 0.35, 'AI-Driven Transformation in Banking GRC  |  Narentheran S  |  Viva voce',
         size=10, color=GREY)
    if number:
        COUNTER['n'] += 1
    label = key.split()[-1] if key else str(COUNTER['n'])
    text(s, 11.9, 7.0, 0.8, 0.35, label, size=10, color=GREY, align=PP_ALIGN.RIGHT)
    k = key or f'Slide {COUNTER["n"]}'
    if k in NOTES:
        s.notes_slide.notes_text_frame.text = NOTES[k]
    return s


def arrow(s, x, y, w=0.45, h=0.3, color=GREY):
    box(s, x, y, w, h, '', fill=color, line=None, shape=MSO_SHAPE.RIGHT_ARROW)


# 1 Title
s = slide(None)
text(s, 0.9, 1.6, 11.5, 2.0, 'AI-Driven Transformation in Banking GRC:\nImplications for Workforce '
     'Management in UK and EU Banks', size=36, bold=True)
box(s, 0.95, 3.75, 2.2, 0.07, '', fill=RED, line=None)
text(s, 0.9, 4.1, 11.5, 2.4, ['**Narentheran S**  |  Roll No. 2406OBA0034  |  Reg. No. 69124100089',
                              'Online MBA, Centre for Distance and Online Education, Anna University, Chennai',
                              'Guide: Dr. RAMYA J, Associate Professor & Head, Department of Management Studies',
                              'Viva voce  |  October 2026'], size=18, color=GREY)

# 2 Why
s = slide('AI is entering bank risk and compliance faster than people can govern it')
for i, (num, lab, red) in enumerate([('75%', 'of UK financial firms\nalready use AI', False),
                                     ('46%', 'have only a partial understanding\nof the AI they use', True),
                                     ('57%', 'say talent and skills\nare a constraint', False)]):
    x = 0.8 + i * 4.0
    text(s, x, 2.0, 3.6, 1.5, num, size=80, bold=True, color=RED if red else BLACK, align=PP_ALIGN.CENTER)
    text(s, x, 3.7, 3.6, 1.2, lab, size=20, color=BLACK if red else GREY, align=PP_ALIGN.CENTER)
text(s, 0.8, 5.5, 11.8, 0.8, 'Adoption is not the problem. **Understanding and skills are.**', size=24,
     align=PP_ALIGN.CENTER)
text(s, 0.8, 6.35, 11.8, 0.4, 'Source: Bank of England & FCA (2024), Artificial intelligence in UK financial services',
     size=11, color=GREY, align=PP_ALIGN.CENTER)

# 3 Management problem
s = slide('This is a management problem, and regulation makes it urgent')
text(s, 0.6, 1.6, 6.5, 0.5, 'Four decisions managers now face', size=18, bold=True, color=GREY)
for i, t in enumerate(['Which GRC roles to **redesign, reduce or create**',
                       'Whether to **build, buy or borrow** skills',
                       'How to **shape the team**', 'How to get staff to **trust AI appropriately**']):
    box(s, 0.6 + (i % 2) * 3.25, 2.2 + (i // 2) * 1.75, 3.05, 1.55, t, line=GREY, size=17)
text(s, 7.5, 1.6, 5.2, 0.5, 'Regulation', size=18, bold=True, color=GREY)
for i, (r, d) in enumerate([('DORA', 'board & staff ICT competence'), ('EU AI Act', 'AI literacy, human oversight'),
                            ('PRA SS1/23', 'skilled model validators'), ('SM&CR', 'named personal accountability')]):
    box(s, 7.5, 2.2 + i * 0.72, 1.7, 0.6, f'**{r}**', fill=LGREY, line=None, size=15)
    text(s, 9.3, 2.27 + i * 0.72, 3.4, 0.55, d, size=15, color=GREY)
box(s, 7.5, 5.2, 5.2, 0.95, 'Competence and human oversight are now legal obligations', fill=RED, line=None,
    size=17, color=WHITE, bold=True)

# 4 Objectives
s = slide('Objectives: build the evidence, then turn it into a framework')
text(s, 0.6, 1.65, 12, 1.0, '**Primary:** examine how AI in GRC is reshaping workforce decisions (role structure, '
     'skills, staffing, employee adaptation) and develop a practical framework for management.', size=19)
for i, (n, t) in enumerate([('i', 'Which GRC activities are automated or augmented, and the effect on role demand'),
                            ('ii', 'How DORA, the AI Act, SS1/23 and SM&CR shape oversight and competence'),
                            ('iii', 'Emerging skill gaps and reskilling strategies'),
                            ('iv', 'Employee responses: usefulness, trust, replaceability'),
                            ('v', 'An integrated workforce transformation framework and recommendations')]):
    y = 2.85 + i * 0.78
    box(s, 0.7, y, 0.6, 0.6, f'**{n}**', fill=RED if n == 'v' else LGREY, line=None, size=16,
        color=WHITE if n == 'v' else BLACK, shape=MSO_SHAPE.OVAL)
    text(s, 1.5, y + 0.07, 11, 0.6, f'[[{t}]]' if n == 'v' else t, size=18)

# 5 Literature
s = slide('Three bodies of research, and none of them looks at the GRC workforce')
for (cx, cy, title, refs) in [(3.2, 1.75, 'Technology acceptance & trust', 'Davis (1989)\nGlikson & Woolley (2020)'),
                              (6.6, 1.75, 'Automation & work', 'Autor (2015); Frey & Osborne (2017)\n'
                               'Cascio & Montealegre (2016)\nBrougham & Haar (2018)'),
                              (4.9, 3.95, 'RegTech', 'Arner et al. (2017)\nBuckley et al. (2020)')]:
    box(s, cx, cy, 3.4, 3.0, '', line=GREY, shape=MSO_SHAPE.OVAL, lw=2.5)
box(s, 3.35, 2.25, 2.4, 1.3, '**Technology acceptance & trust**\nDavis (1989)\nGlikson & Woolley (2020)', line=None, size=12)
box(s, 7.45, 2.15, 2.4, 1.5, '**Automation & work**\nAutor (2015)\nFrey & Osborne (2017)\nCascio & Montealegre (2016)\n'
    'Brougham & Haar (2018)', line=None, size=12)
box(s, 5.4, 5.45, 2.4, 1.2, '**RegTech**\nArner et al. (2017)\nBuckley et al. (2020)', line=None, size=12)
box(s, 5.55, 3.55, 1.95, 1.3, 'GAP', fill=RED, line=None, size=20, color=WHITE, bold=True, shape=MSO_SHAPE.OVAL)
text(s, 10.3, 3.2, 2.8, 2.2, ['[[Research gap]]', 'Workforce management inside regulated banking GRC: '
                              'roles, skills, staffing and adaptation'], size=16)

# 6 Propositions
s = slide('The literature gave four propositions to test')
table(s, 0.7, 1.75, 11.9, [['', 'Proposition', 'Grounded in'],
                           ['P1', 'Routine GRC tasks are automated first; judgement work is augmented', 'Autor (2015); Frey & Osborne (2017)'],
                           ['P2', 'Regulation turns competence and oversight into obligations', 'Buckley et al. (2020); DORA, AI Act, SS1/23, SM&CR'],
                           ['P3', 'AI needs a hybrid technical + regulatory skill profile', 'Cascio & Montealegre (2016)'],
                           ['P4', 'Losing expert knowledge is a hidden restructuring risk   [[EMERGENT]]', 'Interviews (extends Autor, 2015)']],
      [0.8, 6.6, 4.5], size=17, row_h=0.85)

# 7 Design
s = slide('A sequential exploratory mixed-methods design')
for i, (h, b, red) in enumerate([('1  LISTEN', 'Interviews + questionnaire\nThemes T1–T10', False),
                                 ('2  TEST', 'Hypotheses H5–H10\non public data D1–D4', True),
                                 ('3  COMPARE', 'Triangulation\nP1–P4 verdicts + framework', False)]):
    box(s, 0.7 + i * 4.1, 2.0, 4.0, 2.2, f'**{h}**\n{b}', fill=RED if red else LGREY, line=None, size=19,
        color=WHITE if red else BLACK, shape=MSO_SHAPE.CHEVRON)
text(s, 0.9, 4.75, 11.6, 1.6, ['The interviews decided **what** to test. The public data showed **whether it holds** across the industry.',
                               '[[The public datasets are independent sources; they were not produced from the interviews.]]'],
     size=19, align=PP_ALIGN.CENTER, gap=14)

# 8 Evidence
s = slide('Two kinds of evidence: practitioners, and four public datasets')
text(s, 0.6, 1.6, 4.6, 0.5, 'Primary', size=20, bold=True, color=GREY)
text(s, 0.6, 2.15, 4.6, 4.5, ['**Written semi-structured interviews:** 11 questions, 5 sections',
                              '**Questionnaire:** 25 Likert items, 7 constructs adapted from Davis; Glikson & Woolley; Brougham & Haar',
                              '**Sampling:** purposive + snowball; GRC professionals exposed to AI'], size=16, gap=14)
text(s, 5.6, 1.6, 7, 0.5, 'Secondary (public)', size=20, bold=True, color=GREY)
for i, (code, src, cov) in enumerate([('D1', 'Eurostat', 'AI adoption, 27 EU countries, 2021–2025'),
                                      ('D2', 'Bank annual reports', 'Large UK/EU banks, FY2022 vs FY2025; 19 paired'),
                                      ('D3', 'Reed job adverts', '163 UK financial-services GRC adverts'),
                                      ('D4', 'Published surveys', 'BoE/FCA 2022 & 2024, IIA, Thomson Reuters')]):
    y = 2.15 + i * 1.12
    box(s, 5.6, y, 0.95, 0.9, f'**{code}**', fill=LGREY, line=None, size=18)
    text(s, 6.7, y, 6.1, 0.95, [f'**{src}**', (cov, 15)], size=16, gap=2)

# 9 Hypotheses mapping
s = slide('Every hypothesis traces back to something practitioners said')
table(s, 0.5, 1.6, 12.3, [['Interview theme', 'Questionnaire signal', 'Public-data hypothesis'],
                          ['T1  AI spreading, still pilot-heavy', 'Role will change: 4.33', 'H5a, H8: adoption is rising'],
                          ['T2/T3  Routine work absorbed', 'Usefulness: 3.58', 'H10a AI mentions; H6 skill mix in adverts'],
                          ['T4/T5  Humans must challenge AI', 'Can defend AI decision: 4.33', 'H10b: AI governance bodies'],
                          ['T6  Regulation felt unevenly', 'Oversight clear: 3.67', 'H10a regulation mentions; H10d UK vs EU'],
                          ['T7  Hybrid skills', 'Data skills more important: 4.00', 'H5b, H6, H7: skills and pay'],
                          ['T8/T9  Training lags, knowledge at risk', 'Training adequate: 3.33', 'H9 skills constraint; H10c reskilling vs AI'],
                          ['T10  Anxiety, weak change management', 'Replaceability concern: 3.89', 'Not testable on public data']],
      [4.3, 3.6, 4.4], size=15, row_h=0.6, red_rows=(6,))

# 10 Test choice
s = slide('Each test was chosen by the shape of the data')
rows = [('Same countries/banks measured twice', 'Paired t-test + Wilcoxon;  Friedman (4 yrs)', 'H5a, H10a'),
        ('Yes/no for the same banks, two years', 'Exact McNemar', 'H10b'),
        ('Two independent, skewed groups', 'Mann–Whitney U', 'H7, H10d'),
        ('Yes/no skill flags across two groups', 'Chi-square / Fisher + Holm', 'H6'),
        ('Only published percentages', 'Two-proportion & one-sample z', 'H8, H9'),
        ('Strength of a link', 'Pearson / Spearman / OLS', 'H5b, H10c')]
for i, (cond, test, hyp) in enumerate(rows):
    y = 1.65 + i * 0.75
    box(s, 0.6, y, 4.9, 0.6, cond, line=GREY, size=15, align=PP_ALIGN.LEFT)
    arrow(s, 5.62, y + 0.15)
    box(s, 6.2, y, 4.6, 0.6, f'**{test}**', fill=LGREY, line=None, size=15)
    text(s, 10.95, y + 0.1, 1.9, 0.5, hyp, size=15, color=GREY)
box(s, 0.6, 6.25, 12.2, 0.6, 'Questionnaire: descriptive only (means/SD vs neutral 3); significance tests are not meaningful on a small purposive sample',
    fill=None, line=RED, size=15, color=RED, bold=True, lw=2)

# 11 Practitioner quotes
s = slide('Practitioners: routine work goes; judgement and knowledge matter more')
for i, q in enumerate(['"GRC testing roles have disappeared."',
                       '"C-level execs really do not understand what the AI is, and thus are likely to simply trust it."',
                       '"Too much business knowledge exists in people\'s heads… I would prioritise knowledge capture '
                       'and transfer before any significant workforce reductions."']):
    x = 0.6 + i * 4.15
    box(s, x, 1.75, 0.08, 2.9, '', fill=RED, line=None)
    text(s, x + 0.2, 1.75, 3.7, 2.9, q, size=18, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 5.0, 12.2, 1.8, ['**Ten themes by objective:**  (i) automation precedes AI · routine absorbed · roles recomposed   '
                              '(ii) expert review · accountability gap · regulation uneven   (iii) hybrid skills · '
                              '[[knowledge loss]] · build-buy-borrow   (iv) ease welcomed, anxiety, weak change management'],
     size=15, color=GREY)

# 12 Questionnaire
s = slide('Questionnaire: easy to use but hard to see into, and training lags')
img(s, 'q_constructs.png', 0.4, 1.6, w=6.6)
img(s, 'q_gaps.png', 7.1, 1.8, w=5.8)
text(s, 7.2, 5.3, 5.7, 1.4, ['[[Usability–transparency gap]]: 4.33 vs 3.33', '[[Training gap]]: skills matter more (4.00), training adequate 3.33',
                             ('Descriptive means only; not generalised', 12)], size=15, gap=6)

# 13 Adoption
s = slide('AI adoption is surging everywhere  (H5a, H8, H10a supported)')
img(s, 'd1_adoption.png', 0.4, 1.6, w=6.1)
img(s, 'd2_slope.png', 6.9, 1.55, h=4.3)
text(s, 0.6, 5.85, 12.2, 1.0, ['EU large firms **27.5% → 53.6%**, up in 27/27 countries (p < 0.0001)   |   UK FS firms **58% → 75%** (p = 0.008)   |   '
                               'Bank AI mentions **≈10×**, up in [[19 of 19 banks]] (p < 0.0001)'], size=15, align=PP_ALIGN.CENTER)

# 14 Governance
s = slide('Governance structures are growing faster than governance practice')
img(s, 'd2_govbody.png', 0.6, 1.6, h=4.4)
text(s, 5.6, 1.8, 7.2, 4.6, ['**Banks naming an AI governance body: 5% → 42%**  (7 added, 0 dropped; exact McNemar p = 0.016)',
                             '**AI Act / DORA / SS1/23 mentions:** near zero → routine; 18 of 19 banks up (p < 0.0001)',
                             '[[Broader AI-governance wording flat (p ≈ 0.07)]]',
                             '**UK > EU** on AI and AI-governance mentions (Holm p = 0.007 and 0.0004)'], size=18, gap=18)

# 15 Skills
s = slide('Skills are shifting, but reskilling is not keeping up')
img(s, 'd3_skills.png', 0.4, 1.6, w=7.0)
text(s, 7.7, 1.8, 5.2, 4.8, ['AI-exposed adverts want **more data, automation and cyber** skills but **fewer regulatory** skills (all Holm p < 0.05): the hybrid gap',
                             'AI-titled roles: median **£90k vs £60k** (p = 0.009; small group)',
                             '[[Reskilling mentions not linked to AI intensity (ρ = −0.16, p = 0.49)]]'], size=17, gap=18)

# 16 Scorecard
s = slide('Hypothesis scorecard')
sc = [['Hypothesis', 'Verdict'],
      ['H1–H3  Useful, easy to use, roles will change (questionnaire)', '●  Supported (descriptive)'],
      ['H4  Trust is conditional', '◐  Partly'],
      ['H5a, H5b  EU adoption rising; ICT skills linked', '●  Supported'],
      ['H6  Skill mix differs in AI adverts', '●  Supported'],
      ['H7  AI roles pay more', '◐  Partly (AI-titled only)'],
      ['H8  UK FS adoption rose', '●  Supported'],
      ['H9  Skills a majority constraint', '◐  Partly'],
      ['H10a  AI, reskilling, regulation mentions rose', '●  Supported'],
      ['H10a  AI-governance wording rose', '○  Not supported'],
      ['H10b  More AI governance bodies', '●  Supported'],
      ['H10c  AI and reskilling linked', '○  Not supported'],
      ['H10d  UK and EU differ', '◐  Partly (AI, AI governance)']]
table(s, 1.2, 1.55, 10.9, sc, [7.4, 3.5], size=14, row_h=0.41, red_rows=(9, 11))

# 17 Triangulation
s = slide('Triangulation: all four propositions hold')
table(s, 0.5, 1.7, 12.3, [['', 'Interviews', 'Questionnaire', 'Public data', 'Verdict'],
                          ['P1 Routine automated first', 'T1–T3', 'Role change 4.33', 'H5a, H8, H10a', 'Supported'],
                          ['P2 Regulation → obligations', 'T4–T6', 'Items 26–27', 'H10a reg, H10b, H10d', 'Texts & structures; partly in practice'],
                          ['P3 Hybrid skills', 'T7, T9', 'Items 24–25', 'H6, H5b, H9', 'Supported'],
                          ['P4 Knowledge-loss risk', 'T8', '–', 'Indirect (H10c)', 'Qualitative; emergent contribution']],
      [3.0, 1.6, 2.2, 2.7, 2.8], size=15, row_h=0.85, red_rows=(4,))

# 18 Role-impact map
s = slide('AI changes the composition of GRC work, not just its volume')
quad = [('AUTOMATED', 'Data consolidation · reconciliation · MI · first-level alert triage · periodic KYC refresh', LGREY, False),
        ('AUGMENTED', 'Alert investigation · KYC/CDD judgement · model validation · audit fieldwork · regulatory change', LGREY, False),
        ('NEWLY CREATED', 'AI governance lead · AI model-risk validator · AI audit lead · human-in-the-loop supervisor of AI agents', LGREY, False),
        ('HUMAN-RETAINED', 'Interpreting regulatory intent · challenging decisions · assurance to boards · final accountability (SM&CR, DORA)', None, True)]
for i, (h, b, fill, red) in enumerate(quad):
    x, y = 0.7 + (i % 2) * 6.1, 1.65 + (i // 2) * 2.6
    box(s, x, y, 5.9, 2.4, f'**{h}**\n{b}', fill=fill, line=RED if red else None, size=16,
        color=RED if red else BLACK, lw=3)

# 19 Framework
s = slide('A seven-stage framework: protect knowledge before restructuring')
stages = ['Diagnose', 'Protect knowledge', 'Redesign roles', 'Build oversight competence', 'Reskill & source',
          'Secure trust', 'Lead change & retain']
for i, st in enumerate(stages):
    red = i == 1
    x = 0.55 + i * 1.78
    box(s, x, 2.1, 1.45, 1.75, f'**{i + 1}**\n{st}', fill=RED if red else LGREY, line=None, size=14,
        color=WHITE if red else BLACK)
    if i < 6:
        arrow(s, x + 1.47, 2.83, w=0.28, h=0.3)
text(s, 0.6, 4.25, 12.2, 1.2, '[[Why this order?]] The experts whose routine work is automated are the same people needed to validate the AI.',
     size=19, align=PP_ALIGN.CENTER)
box(s, 0.6, 5.6, 12.2, 0.7, 'Enablers: data quality & lineage  ·  clear accountability map  ·  workforce indicators tracked alongside AI indicators',
    line=GREY, size=15, color=GREY)

# 20 Recommendations
s = slide('Recommendations, conclusion and limits')
for i, (h, b) in enumerate([('GRC leaders', 'Map activities, not job titles\nCapture knowledge before cutting roles\nMake oversight operational'),
                            ('HR', 'Close the training gap: role-tiered AI literacy\nBuild, buy, borrow deliberately\nCareer paths for junior staff'),
                            ('Boards & senior managers', 'Close the AI-understanding gap at the top\nPair every AI plan with a people plan\nTrack workforce indicators')]):
    x = 0.6 + i * 4.1
    text(s, x, 1.65, 3.9, 0.5, h, size=19, bold=True, color=RED)
    text(s, x, 2.2, 3.9, 2.2, b.split('\n'), size=16, gap=8)
box(s, 0.6, 4.55, 12.2, 0.85, '**Conclusion:** AI in GRC succeeds only if the people who answer for it can understand and challenge it.',
    fill=LGREY, line=None, size=18)
text(s, 0.6, 5.6, 12.2, 1.2, ['**Limits:** small purposive primary sample (analytical generalisation) · disclosure ≠ action · 19 paired banks · one UK job board',
                              '**Next:** larger surveys · annual tracking of reports and adverts · EU job data · in-bank case studies'], size=14, color=GREY)

# 21 Thank you
s = slide(None)
text(s, 0.9, 2.3, 11.5, 1.2, 'Thank you', size=54, bold=True, align=PP_ALIGN.CENTER)
box(s, 5.9, 3.55, 1.5, 0.07, '', fill=RED, line=None)
text(s, 0.9, 3.9, 11.5, 1.0, 'Questions welcome', size=24, color=GREY, align=PP_ALIGN.CENTER)
text(s, 0.9, 5.0, 11.5, 0.8, '[[AI in GRC succeeds only if the people who answer for it can understand and challenge it.]]',
     size=18, align=PP_ALIGN.CENTER)

# Backup slides
s = slide('Backup: D2 bank sample and coding', key='Backup B1', number=False)
text(s, 0.6, 1.65, 12.2, 5.2, ['**Both years (19):** ABN AMRO, BBVA, Banco Santander, Bank of Ireland, Barclays, Danske Bank, Deutsche Bank, Erste Group, '
                               'Groupe BPCE, HSBC, Intesa Sanpaolo, KBC, NatWest, Nationwide, Nordea, Santander UK, Société Générale, Standard Chartered, UniCredit',
                               '**FY2025 only:** Handelsbanken, ING, SEB   |   **FY2022 only:** Rabobank',
                               '**Not coded:** AIB, BNP Paribas, CaixaBank, Commerzbank, Crédit Agricole, Lloyds (no full English report / access blocked)',
                               '**Coding:** fixed keyword lists for AI, reskilling, AI governance and regulation; count ÷ total words × 10,000; '
                               'binary flag for a named AI governance body. Each file checked against its title page.'], size=15, gap=12)
s = slide('Backup: full H10 results', key='Backup B2', number=False)
table(s, 0.5, 1.55, 12.3, [['Test', 'Statistic', 'p', 'Effect', 'n'],
                           ['AI density, paired t / Wilcoxon', 't = 6.61 / W = 190', '<0.0001', '+2.61; 19/19 up', '19'],
                           ['Reskilling density', 't = 2.09 / W = 142', '0.026 / 0.030', '+0.16; 13/19 up', '19'],
                           ['AI-governance density', 't = 1.56 / W = 132', '0.068 / 0.072', '+0.15; 13/19 up', '19'],
                           ['Regulation density', 't = 5.42 / W = 171', '<0.0001', '+0.20; 18/19 up', '19'],
                           ['AI governance body, exact McNemar', '–', '0.016', '7 up, 0 down', '19'],
                           ['AI vs reskilling, Spearman (FY2025)', 'ρ = −0.16', '0.487', '', '22'],
                           ['UK vs EU AI density, Mann–Whitney', 'U = 87', 'Holm 0.007', '3.96 vs 2.00', '6+16'],
                           ['UK vs EU reskilling', 'U = 23', 'Holm 0.140', '0.22 vs 0.42', '6+16'],
                           ['UK vs EU AI governance', 'U = 94', 'Holm 0.0004', '2.01 vs 0.50', '6+16'],
                           ['UK vs EU regulation', 'U = 45', 'Holm 0.858', '0.21 vs 0.15', '6+16']],
      [4.2, 2.4, 1.9, 2.6, 1.2], size=13, row_h=0.45)
s = slide('Backup: questionnaire item means (1–5)', key='Backup B3', number=False)
items = [('7 Improves performance', '4.00'), ('8 Increases volume', '3.33'), ('9 Improves quality', '3.33'), ('10 Useful overall', '3.67'),
         ('11 Easy to use', '4.33'), ('12 Understand what it does', '3.33'), ('13 Easy to learn', '4.33'), ('14 Understand outputs', '3.67'),
         ('15 Outputs reliable', '3.67'), ('16 Could defend decision', '4.33'), ('17 Rely without re-performing', '3.67'), ('18 Can override', '4.00'),
         ('19 Role could be done by AI', '3.33'), ('20 Role will change', '4.33'), ('21 Career concern', '4.00'), ('22 Skills sufficient', '3.67'),
         ('23 Training adequate', '3.33'), ('24 Data skills more important', '4.00'), ('25 Judgement more important', '4.00'), ('26 Accountable people knowledgeable', '4.00'),
         ('27 Oversight clear in practice', '3.67'), ('28 Management communicated', '4.00'), ('29 Positive about AI', '4.00'), ('30 Would consider leaving', '3.67'),
         ('31 Committed to organisation', '4.00')]
half = 13
for col in range(2):
    part = items[col * half:(col + 1) * half]
    table(s, 0.6 + col * 6.2, 1.55, 5.9, [['Item', 'Mean']] + [list(p) for p in part], [4.9, 1.0], size=12, row_h=0.37,
          red_rows=tuple(i + 1 for i, p in enumerate(part) if p[0].split()[0] in ('12', '17', '23')))
s = slide('Backup: regulation mapped to workforce obligations', key='Backup B4', number=False)
table(s, 0.5, 1.6, 12.3, [['Instrument', 'Workforce requirement', 'Implication'],
                          ['DORA (2022/2554)', 'Board ICT knowledge (Art. 5); training for all staff (Art. 13)', 'Capability becomes auditable; training records = evidence'],
                          ['EU AI Act (2024/1689) + Omnibus', 'AI literacy (Art. 4); effective human oversight (Art. 14)', 'Role-based AI literacy; trained overseers'],
                          ['PRA SS1/23', 'Accountable SMF; documented lifecycle roles and validator skills', 'AI-capable model validation roles'],
                          ['SM&CR', 'Existing senior managers accountable for AI use', 'AI understanding gap = personal exposure'],
                          ['AMLR / AMLA (2024)', 'Compliance manager + AML compliance officer', 'Senior AML roles own AI-enabled monitoring']],
      [3.0, 4.9, 4.4], size=14, row_h=0.75)
s = slide('Backup: why the questionnaire is descriptive and public data carry the tests', key='Backup B5', number=False)
text(s, 0.6, 1.7, 12.2, 5.0, ['**Questionnaire:** small purposive sample of senior practitioners, used to validate interview themes. Significance tests and '
                              "Cronbach's alpha are not meaningful at this size, so means/SD are compared with the neutral point 3.",
                              '**Interviews:** free-text data, so thematic analysis is the right method.',
                              '**Public data:** large, independent samples (27 countries, 19 paired banks, 163 adverts, surveys of 71–6,506) '
                              'allow proper inferential tests.',
                              '[[The interviews decided what to test; the public data tested whether it holds.]]'], size=18, gap=16)

os.makedirs(OUT, exist_ok=True)
path = os.path.join(OUT, 'Viva_Presentation_Narentheran_S.pptx')
prs.save(path)
print('saved', path, len(prs.slides), 'slides; notes on', sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()))
