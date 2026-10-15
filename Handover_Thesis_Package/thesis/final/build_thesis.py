#!/usr/bin/env python3
"""Build the final thesis .docx/.pdf in the Anna University CDOE MBA project-report format.

Reads the Markdown chapters in this folder and writes out/Thesis_Narentheran_S.docx (+ .pdf when
LibreOffice is available). Contents, tables and figures lists are real Word TOC fields, pre-filled
with page numbers taken from a LibreOffice rendering; Word refreshes them on open.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(HERE, '..', 'stats', 'out')
LOGO = os.path.join(HERE, 'assets', 'anna_university_logo.png')
OUT_DIR = os.path.join(HERE, 'out')
OUT_NAME = 'Thesis_Narentheran_S'

TITLE = ('AI-DRIVEN TRANSFORMATION IN BANKING GRC: IMPLICATIONS FOR WORKFORCE '
         'MANAGEMENT IN UK AND EU BANKS')
STUDENT = 'NARENTHERAN S'
ROLL = '2406OBA0034'
REG = '69124100089'

MARGIN = dict(top=32, bottom=27, left=38, right=22)  # mm, within the prescribed ranges
TEXT_W = 210 - MARGIN['left'] - MARGIN['right']
HEADING_DROP = Pt(51)  # top margin 32 mm + 18 mm = heading 50 mm below the top edge
J, C, L, R = (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER,
              WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT)

BODY_FILES = ['01_chapter1_introduction.md', '02_chapter2_literature.md',
              '03_chapter3_methodology.md', '04_chapter4_analysis.md',
              '05_chapter5_conclusions.md', '06_appendices.md', '07_references.md']


# ---------------------------------------------------------------- low-level helpers
def strip_theme_fonts(rpr_parent):
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        rfonts.attrib.pop(qn(a), None)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(a), 'Times New Roman')


def style_font(style, size=12, bold=None, italic=None):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = None
    strip_theme_fonts(style.element)
    color = style.element.rPr.find(qn('w:color'))
    if color is not None:
        style.element.rPr.remove(color)


def add_runs(p, text, size=None, bold=False, italic=False):
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[A-Za-z(][^*]*?[^\s*]\*)', text):
        if not tok:
            continue
        b, i, t = bold, italic, tok
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            b, t = True, tok[2:-2]
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            i, t = True, tok[1:-1]
        r = p.add_run(t.replace('`', ''))
        if b:
            r.bold = True
        if i:
            r.italic = True
        if size:
            r.font.size = Pt(size)
    return p


def fld_char(kind):
    r = OxmlElement('w:r')
    f = OxmlElement('w:fldChar')
    f.set(qn('w:fldCharType'), kind)
    r.append(f)
    return r


def instr(text):
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = text
    r.append(it)
    return r


def simple_field(p, code, result=''):
    p._p.append(fld_char('begin'))
    p._p.append(instr(f' {code} '))
    p._p.append(fld_char('separate'))
    if result:
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = result
        r.append(t)
        p._p.append(r)
    p._p.append(fld_char('end'))


def tc_field(p, text, level=1):
    p._p.append(fld_char('begin'))
    p._p.append(instr(f' TC "{text}" \\l {level} '))
    p._p.append(fld_char('end'))


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    tcpr.append(sh)


def page_numbering(section, fmt, start):
    sp = section._sectPr
    pg = sp.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sp.append(pg)
    pg.set(qn('w:fmt'), fmt)
    pg.set(qn('w:start'), str(start))


def norm(s):
    return re.sub(r'[^a-z0-9]', '', s.lower())


def roman(n):
    out = ''
    for v, s in ((1000, 'm'), (900, 'cm'), (500, 'd'), (400, 'cd'), (100, 'c'), (90, 'xc'),
                 (50, 'l'), (40, 'xl'), (10, 'x'), (9, 'ix'), (5, 'v'), (4, 'iv'), (1, 'i')):
        while n >= v:
            out += s
            n -= v
    return out


# ---------------------------------------------------------------- document builder
class Thesis:
    def __init__(self, pages=None):
        self.doc = Document()
        self.pages = pages
        self.entries = []          # dicts: kind (toc/lot/lof), level, text
        self.placeholders = {}
        self.fresh_page = True     # next heading needs no page break
        self._setup()

    # styles and page setup
    def _setup(self):
        d = self.doc
        st = d.styles
        normal = st['Normal']
        style_font(normal, 12)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.widow_control = True

        def new(name, base='Normal'):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st[base]
            style_font(s, 12)
            return s

        body = new('Body Text Thesis')
        body.paragraph_format.alignment = J
        body.paragraph_format.first_line_indent = Mm(20)
        body.paragraph_format.space_after = Pt(12)

        for name, caps in (('Heading 2', True), ('Heading 3', False)):
            h = st[name]
            style_font(h, 12, bold=True, italic=False)
            h.font.all_caps = False
            hp = h.paragraph_format
            hp.alignment = L
            hp.space_before = Pt(18)
            hp.space_after = Pt(12)
            hp.keep_with_next = True
            hp.line_spacing = 1.5

        tt = new('Table Text Thesis')
        style_font(tt, 10.5)
        tt.paragraph_format.line_spacing = 1.0
        tt.paragraph_format.alignment = L

        for name, before, after in (('Table Caption', 18, 6), ('Figure Caption', 6, 18)):
            s = new(name)
            style_font(s, 12, bold=True)
            s.paragraph_format.alignment = C
            s.paragraph_format.line_spacing = 1.0
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            if name == 'Table Caption':
                s.paragraph_format.keep_with_next = True

        lst = new('List Thesis')
        lst.paragraph_format.alignment = J
        lst.paragraph_format.left_indent = Mm(28)
        lst.paragraph_format.first_line_indent = Mm(-8)
        lst.paragraph_format.space_after = Pt(6)

        ref = new('Reference Thesis')
        ref.paragraph_format.line_spacing = 1.0
        ref.paragraph_format.alignment = L
        ref.paragraph_format.left_indent = Mm(10)
        ref.paragraph_format.first_line_indent = Mm(-10)
        ref.paragraph_format.space_after = Pt(8)

        for lvl, ind in ((1, 0), (2, 8), (3, 16)):
            s = new(f'TOC {lvl}')
            s.paragraph_format.left_indent = Mm(ind)
            s.paragraph_format.space_after = Pt(0)
            s.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)
        tof = new('Table of Figures')
        tof.paragraph_format.left_indent = Mm(0)
        tof.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT,
                                                    WD_TAB_LEADER.DOTS)

        sec = d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin = Mm(MARGIN['top']), Mm(MARGIN['bottom'])
        sec.left_margin, sec.right_margin = Mm(MARGIN['left']), Mm(MARGIN['right'])
        sec.header_distance, sec.footer_distance = Mm(20), Mm(12)
        sec.different_first_page_header_footer = True
        page_numbering(sec, 'lowerRoman', 1)
        hp = sec.header.paragraphs[0]
        hp.alignment = R
        hp.paragraph_format.line_spacing = 1.0
        simple_field(hp, 'PAGE', '1')

        upd = OxmlElement('w:updateFields')
        upd.set(qn('w:val'), 'true')
        d.settings.element.append(upd)
        d.core_properties.title = TITLE.title()
        d.core_properties.author = 'Narentheran S'

    # basic blocks
    def p(self, text='', style='Body Text Thesis', align=None, size=None, bold=False,
          italic=False, before=None, after=None, spacing=None, indent=None, keep=False):
        para = self.doc.add_paragraph(style=style)
        if align is not None:
            para.alignment = align
        pf = para.paragraph_format
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if spacing is not None:
            pf.line_spacing = spacing
        if indent is not None:
            pf.first_line_indent = Mm(indent)
        if keep:
            pf.keep_with_next = True
        add_runs(para, text, size=size, bold=bold, italic=italic)
        self.fresh_page = False
        return para

    def page_heading(self, text, toc_text=None, label=None):
        """Centred heading 50 mm below the top edge, starting a new page."""
        first = self.p(label or text, style='Normal', align=C, size=14, bold=True,
                       before=HEADING_DROP.pt, after=0, keep=True)
        if not self._take_fresh():
            first.paragraph_format.page_break_before = True
        last = first
        if label:
            last = self.p(text, style='Normal', align=C, size=14, bold=True, before=24, keep=True)
        last.paragraph_format.space_after = Pt(36)
        if toc_text:
            tc_field(last, toc_text, 1)
            self.entries.append(dict(kind='toc', level=1, text=toc_text, start=True))
        return last

    def _take_fresh(self):
        fresh = getattr(self, '_fresh_flag', True)
        self._fresh_flag = False
        return fresh

    def placeholder(self, key):
        para = self.doc.add_paragraph(style='Normal')
        self.placeholders[key] = para

    def heading(self, text, level):
        para = self.doc.add_heading(level=level)
        add_runs(para, text)
        self.entries.append(dict(kind='toc', level=level, text=text))
        return para

    def caption(self, text, kind):
        style = 'Table Caption' if kind == 'lot' else 'Figure Caption'
        para = self.p(text, style=style)
        self.entries.append(dict(kind=kind, level=1, text=text))
        return para

    def table(self, rows, size=10.5, header=True, borders=True):
        ncol = max(len(r) for r in rows)
        lens = [max(6, min(70, max(len(re.sub(r'\*', '', r[i])) if i < len(r) else 0
                                     for r in rows))) for i in range(ncol)]
        weights = [ln ** 0.75 for ln in lens]
        widths = [w / sum(weights) * TEXT_W for w in weights]
        min_w = [max(16 if ncol > 2 else 25,
                     3 + 2.2 * max((len(word) for r in rows if i < len(r)
                                    for word in re.sub(r'\*', '', r[i]).split()), default=0))
                 for i in range(ncol)]
        widths = [max(w, m) for w, m in zip(widths, min_w)]
        widths = [w / sum(widths) * TEXT_W for w in widths]
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = 'Table Grid' if borders else None
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncol):
                txt = row[ci] if ci < len(row) else ''
                cell = cells[ci]
                cell.width = Mm(widths[ci])
                para = cell.paragraphs[0]
                para.style = self.doc.styles['Table Text Thesis']
                add_runs(para, txt, size=size, bold=(header and ri == 0))
                if header and ri == 0 and borders:
                    shade(cell, 'D9D9D9')
            trpr = t.rows[ri]._tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit')
            cs.set(qn('w:val'), 'true')
            trpr.append(cs)
            if header and ri == 0:
                th = OxmlElement('w:tblHeader')
                th.set(qn('w:val'), 'true')
                trpr.append(th)
        for ci, col in enumerate(t.columns):
            col.width = Mm(widths[ci])
        self.p('', style='Normal', spacing=1.0, after=12, size=6)
        return t

    def figure(self, path):
        with Image.open(path) as im:
            aspect = im.height / im.width
        width = 145
        if width * aspect > 180:
            width = 180 / aspect
        para = self.doc.add_paragraph(style='Normal')
        para.alignment = C
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_before = Pt(12)
        para.add_run().add_picture(path, width=Mm(width))

    # ------------------------------------------------------------ front matter
    def title_page(self):
        def tp(text, bold=False, italic=False, before=0, spacing=1.0):
            return self.p(text, style='Normal', align=C, bold=bold, italic=italic,
                          before=before, after=0, spacing=spacing)
        tp(TITLE, bold=True, before=6, spacing=1.5)
        tp('By', bold=True, before=36)
        tp(STUDENT, bold=True, before=18)
        tp(f'Roll No. {ROLL}', before=18)
        tp(f'Reg. No. {REG}', before=6)
        tp('A PROJECT REPORT', before=30)
        tp('Submitted to the', before=18)
        tp('FACULTY OF MANAGEMENT SCIENCES', bold=True, before=6)
        tp('in partial fulfillment for the award of the degree', italic=True, before=30)
        tp('of', italic=True, before=6)
        tp('MASTER OF BUSINESS ADMINISTRATION', bold=True, before=6)
        logo = self.doc.add_paragraph(style='Normal')
        logo.alignment = C
        logo.paragraph_format.space_before = Pt(24)
        logo.add_run().add_picture(LOGO, width=Mm(26))
        tp('CENTRE FOR DISTANCE AND ONLINE EDUCATION', bold=True, before=18)
        tp('ANNA UNIVERSITY', bold=True)
        tp('CHENNAI 600 025', bold=True)
        tp('[MONTH, YEAR]', before=36)
        self._fresh_flag = False

    def bonafide(self):
        self.page_heading('BONAFIDE CERTIFICATE')
        self.p(f'Certified that the Project report titled **"{TITLE}"** is the bonafide work of '
               f'**Mr. {STUDENT}** who carried out the work under my supervision. Certified further '
               'that to the best of my knowledge the work reported herein does not form part of any '
               'other project report or dissertation on the basis of which a degree or award was '
               'conferred on an earlier occasion on this or any other candidate.', spacing=2.0)
        self.p('', style='Normal', before=36)
        self.table([
            ['', ''],
            ['**Signature of Student**', '**Signature of Project supervisor / Guide**'],
            [f'Name: {STUDENT}', 'Name: [GUIDE NAME]'],
            ['(in capital letters)', '(in capital letters)'],
            [f'Roll No.: {ROLL}', 'Designation: [GUIDE DESIGNATION]'],
            [f'Reg. No.: {REG}', 'Address: [GUIDE ADDRESS]'],
        ], size=12, header=False, borders=False)
        self.p('Submitted for the Viva-Voce Examination held on ____________________',
               style='Normal', align=L, before=36)
        self.p('**Signature of External Examiner**', style='Normal', align=R, before=60)

    def viva_placeholder(self):
        self.page_heading('CERTIFICATE OF VIVA-VOCE EXAMINATION')
        self.p('[PLACEHOLDER: VIVA-VOCE CERTIFICATE]', style='Normal', align=C, bold=True,
               before=120)
        self.p('Insert the signed Certificate of Viva-Voce Examination (Annexure III) on this page.',
               style='Normal', align=C, italic=True, before=12)

    def abstract(self, paras):
        self.page_heading('ABSTRACT', toc_text='ABSTRACT')
        for t in paras:
            self.p(t, spacing=2.0, after=0)
        self.p('(Signature of the Candidate)', style='Normal', align=R, before=12)
        self.p(STUDENT, style='Normal', align=R, bold=True)

    def acknowledgement(self, paras):
        self.page_heading('ACKNOWLEDGEMENT')
        for t in paras:
            self.p(t, spacing=2.0, after=0)
        self.p('(Signature)', style='Normal', align=R, before=36)
        self.p(STUDENT, style='Normal', align=R, bold=True)

    def contents_pages(self, abbreviations):
        self.page_heading('TABLE OF CONTENTS')
        self.p('**PAGE NO.**', style='Normal', align=R, after=6)
        self.placeholder('toc')
        self.page_heading('LIST OF TABLES', toc_text='LIST OF TABLES')
        self.p('**PAGE NO.**', style='Normal', align=R, after=6)
        self.placeholder('lot')
        self.page_heading('LIST OF FIGURES', toc_text='LIST OF FIGURES')
        self.p('**PAGE NO.**', style='Normal', align=R, after=6)
        self.placeholder('lof')
        self.page_heading('LIST OF ABBREVIATIONS', toc_text='LIST OF ABBREVIATIONS')
        for abbr, meaning in abbreviations:
            para = self.p(f'**{abbr}**\t{meaning}', style='Normal', align=L)
            para.paragraph_format.tab_stops.add_tab_stop(Mm(32))
            para.paragraph_format.left_indent = Mm(32)
            para.paragraph_format.first_line_indent = Mm(-32)

    # ------------------------------------------------------------ main text
    def markdown(self, text):
        lines = text.split('\n')
        i, buf = 0, []
        pending_figure = False

        def flush():
            if buf:
                self.p(' '.join(s.strip() for s in buf))
                buf.clear()
        in_refs = False
        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            if s.startswith('<!--'):
                while '-->' not in lines[i]:
                    i += 1
                i += 1
                continue
            if not s or s == '---':
                flush()
                i += 1
                continue
            m = re.match(r'^# (CHAPTER|APPENDIX) (\d+) - (.+)$', s)
            if m:
                flush()
                word, num, title = m.groups()
                self.page_heading(title.upper(), toc_text=f'{word} {num} - {title.upper()}',
                                  label=f'{word} {num}')
                i += 1
                continue
            if s == '# REFERENCES':
                flush()
                self.page_heading('REFERENCES', toc_text='REFERENCES')
                in_refs = True
                i += 1
                continue
            if s.startswith('### '):
                flush()
                self.heading(s[4:], 3)
                i += 1
                continue
            if s.startswith('## '):
                flush()
                self.heading(s[3:], 2)
                i += 1
                continue
            if re.match(r'^Table (A )?\d+(\.\d+)+ ', s) and not s.endswith('.'):
                flush()
                self.caption(s, 'lot')
                i += 1
                continue
            if s.startswith('|'):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    row = lines[i].strip()
                    if not re.match(r'^\|?\s*:?-{3,}', row):
                        rows.append([c.strip() for c in row.strip('|').split('|')])
                    i += 1
                self.table(rows)
                continue
            m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)', s)
            if m:
                flush()
                path = os.path.normpath(os.path.join(HERE, m.group(1)))
                self.figure(path)
                pending_figure = True
                i += 1
                continue
            if pending_figure and s.startswith('Figure '):
                self.caption(s, 'lof')
                pending_figure = False
                i += 1
                continue
            if s.startswith('Source:'):
                flush()
                self.p(s, style='Normal', align=L, italic=True, size=10.5, after=12, spacing=1.0)
                i += 1
                continue
            m = re.match(r'^(- |(\d+|[ivx]+)\. )(.*)$', s)
            if m:
                flush()
                marker = '\u2022' if m.group(1) == '- ' else m.group(1).strip()
                if in_refs:
                    self.p(f'{marker}\t{m.group(3)}', style='Reference Thesis')
                else:
                    self.p(f'{marker}\t{m.group(3)}', style='List Thesis')
                i += 1
                continue
            if re.match(r'^\*\*[^*]+\*\*$', s) or re.match(r'^\*[^*]+\*$', s):
                flush()
                self.p(s, style='Normal', align=L, before=6, after=6, keep=True)
                i += 1
                continue
            buf.append(s)
            i += 1
        flush()

    # ------------------------------------------------------------ lists of contents
    def fill_lists(self):
        pages = self.pages
        groups = {
            'toc': ('TOC \\o "2-3" \\l "1-1" \\z', lambda e: f'TOC {e["level"]}'),
            'lot': ('TOC \\z \\t "Table Caption,1"', lambda e: 'Table of Figures'),
            'lof': ('TOC \\z \\t "Figure Caption,1"', lambda e: 'Table of Figures'),
        }
        for kind, (code, style_of) in groups.items():
            ph = self.placeholders[kind]
            items = [(n, e) for n, e in enumerate(self.entries) if e['kind'] == kind]
            paras = []
            for n, e in items:
                page = pages[n] if pages else '00'
                para = ph.insert_paragraph_before(style=style_of(e))
                para.paragraph_format.line_spacing = 1.5
                if kind == 'toc' and e['level'] == 1:
                    para.paragraph_format.space_before = Pt(6)
                add_runs(para, f'{e["text"]}\t{page}', bold=(kind == 'toc' and e['level'] == 1))
                paras.append(para)
            first, last = paras[0], paras[-1]
            ppr = first._p.find(qn('w:pPr'))
            pos = list(first._p).index(ppr) + 1 if ppr is not None else 0
            for el in reversed([fld_char('begin'), instr(f' {code} '), fld_char('separate')]):
                first._p.insert(pos, el)
            last._p.append(fld_char('end'))
            ph._p.getparent().remove(ph._p)

    def save(self, path):
        self.fill_lists()
        self.doc.save(path)


# ---------------------------------------------------------------- assembly
def read(name):
    with open(os.path.join(HERE, name), encoding='utf-8') as f:
        return f.read()


def front_sections():
    text = read('00_front_matter.md')
    parts = dict(re.findall(r'^# ([A-Z ]+)\n(.*?)(?=^# |\Z)', text, flags=re.S | re.M))

    def paras(block):
        out = []
        for chunk in re.split(r'\n\s*\n', block):
            c = chunk.strip()
            if not c or c.startswith('(Signature') or c.startswith('---') or c.startswith('<!--') \
                    or c.strip('*') == STUDENT or c == STUDENT:
                continue
            out.append(' '.join(c.split('\n')))
        return out
    abbrs = re.findall(r'^\| ([^|]+?) \| ([^|]+?) \|$', parts['LIST OF ABBREVIATIONS'], flags=re.M)
    abbrs = [(a, b) for a, b in abbrs if a not in ('Abbreviation',) and not a.startswith('-')]
    return paras(parts['ABSTRACT']), paras(parts['ACKNOWLEDGEMENT']), abbrs


def build(pages, path):
    abstract, ack, abbrs = front_sections()
    t = Thesis(pages)
    t.title_page()
    t.bonafide()
    t.viva_placeholder()
    t.abstract(abstract)
    t.acknowledgement(ack)
    t.contents_pages(abbrs)
    sec = t.doc.add_section(WD_SECTION.NEW_PAGE)
    sec.different_first_page_header_footer = False
    page_numbering(sec, 'decimal', 1)
    t._fresh_flag = True
    for name in BODY_FILES:
        t.markdown(read(name))
    t.save(path)
    return t.entries


def to_pdf(docx_path, outdir):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return None
    subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', outdir, docx_path],
                   check=True, capture_output=True, timeout=600)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')


def page_texts(pdf):
    raw = subprocess.run(['pdftotext', pdf, '-'], check=True, capture_output=True,
                         text=True).stdout
    pages = []
    for pg in raw.split('\f'):
        lines = [ln for ln in pg.strip().split('\n') if ln.strip()]
        if lines and re.fullmatch(r'[ivxlcdm]+|\d+', lines[0].strip()):
            lines = lines[1:]
        pages.append(norm(' '.join(lines)))
    return pages


def locate(entries, texts):
    start_pages = [i for i, t in enumerate(texts)]
    found = [None] * len(entries)
    ptr = {'toc': 0, 'lot': None, 'lof': None}
    body_start = None
    for n, e in enumerate(entries):
        key = norm(e['text'])[:45]
        k = e['kind']
        if k != 'toc':
            continue
        p0 = ptr['toc']
        for p in range(p0, len(texts)):
            hit = texts[p].startswith(key) if e.get('start') else key in texts[p]
            if hit:
                found[n] = p
                ptr['toc'] = p
                break
        if e['text'].startswith('CHAPTER 1 ') and found[n] is not None:
            body_start = found[n]
    for kind in ('lot', 'lof'):
        p0 = body_start or 0
        for n, e in enumerate(entries):
            if e['kind'] != kind:
                continue
            key = norm(e['text'])[:45]
            for p in range(p0, len(texts)):
                if key in texts[p]:
                    found[n] = p
                    p0 = p
                    break
    labels = []
    for p in found:
        if p is None:
            labels.append('?')
        elif body_start is not None and p >= body_start:
            labels.append(str(p - body_start + 1))
        else:
            labels.append(roman(p + 1))
    return labels, found, body_start


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    final = os.path.join(OUT_DIR, OUT_NAME + '.docx')
    with tempfile.TemporaryDirectory() as tmp:
        draft = os.path.join(tmp, OUT_NAME + '.docx')
        entries = build(None, draft)
        pdf = to_pdf(draft, tmp)
        if not pdf:
            shutil.copy(draft, final)
            print('LibreOffice not found: lists keep placeholder page numbers.')
            return
        labels, found, body_start = locate(entries, page_texts(pdf))
        build(labels, final)
    pdf = to_pdf(final, OUT_DIR)
    labels2, found2, body2 = locate(entries, page_texts(pdf))
    appendix = next(f for e, f in zip(entries, found2) if e['text'].startswith('APPENDIX 1 '))
    missing = [e['text'] for e, lab in zip(entries, labels2) if lab == '?']
    print(f'Wrote {final} and {pdf}')
    print(f'Total pages: {len(page_texts(pdf)) - 1}; Chapter 1 to end of Chapter 5: '
          f'{appendix - body2} pages (limit 60)')
    print('Page labels stable between passes:', labels == labels2)
    if missing:
        print('Entries not located:', missing)


if __name__ == '__main__':
    sys.exit(main())
