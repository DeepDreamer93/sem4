from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d=Document()
for s in d.sections:
    s.left_margin=s.right_margin=Cm(2.5); s.top_margin=s.bottom_margin=Cm(2.2)
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
for h in ['Heading 1','Heading 2','Heading 3','Title']:
    d.styles[h].font.name='Times New Roman'; d.styles[h].font.color.rgb=RGBColor(0x1F,0x3A,0x5F)
    d.styles[h].element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')

def P(t,b=False,i=False,al=None):
    p=d.add_paragraph(); r=p.add_run(t); r.bold=b; r.italic=i
    p.paragraph_format.space_after=Pt(6); p.alignment=al if al is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p
def B(items,style='List Bullet'):
    for t in items:
        p=d.add_paragraph(style=style)
        if isinstance(t,tuple):
            r=p.add_run(t[0]); r.bold=True; p.add_run(t[1])
        else: p.add_run(t)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
def Q(t):
    p=d.add_paragraph(); r=p.add_run('\u201c'+t+'\u201d'); r.italic=True
    p.paragraph_format.left_indent=Cm(1); p.paragraph_format.space_after=Pt(6)
def shade(c,hexc):
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)
def T(head,rows,widths=None,cap=None):
    if cap: P(cap,b=True,al=WD_ALIGN_PARAGRAPH.LEFT)
    t=d.add_table(rows=1,cols=len(head)); t.style='Table Grid'
    for i,h in enumerate(head):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(10); shade(c,'D9E2F3')
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=''; r=cs[i].paragraphs[0].add_run(v); r.font.size=Pt(10)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    d.add_paragraph()

# ---------------- Title
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('AI-Driven Transformation in Banking GRC: Implications for Workforce Management in UK and EU Banks'); r.bold=True; r.font.size=Pt(16)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Data Analysis and Interpretation: Primary Findings, Secondary Evidence and Preliminary Framework'); r.font.size=Pt(13); r.italic=True
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Narentheran S  |  Roll No: 2406OBA0034  |  Reg No: 69124100089').font.size=Pt(11)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Working draft prepared to build on the First Review Report').font.size=Pt(10)

# ---------------- 1
d.add_heading('1 Introduction and Analytical Approach',1)
P('This chapter presents the analysis of the evidence gathered for the study. It follows the mixed-methods, qualitatively dominant design set out in the First Review Report and is organised around the five secondary objectives. The evidence is drawn from separate streams that are analysed and then triangulated: (a) written practitioner interviews completed by senior professionals working in or alongside GRC functions serving UK and EU banks; (b) the structured questionnaire built on constructs adapted from Davis (1989), Brougham and Haar (2018) and Glikson and Woolley (2020); and (c) publicly available secondary evidence, comprising UK and EU regulatory and supervisory texts, industry and professional-body research, and the publicly advertised role profiles and skill requirements for GRC positions in the banking labour market.')
P('Consistent with the ethical commitments in the First Review Report, no respondent, employer or client is identified, and no individual bank is named. Secondary evidence on job descriptions and workforce practice is reported at the level of the sector, synthesised from public role profiles, recruiter and professional-body publications and regulatory documents, rather than attributed to any institution. The primary sample was deliberately purposive and small, and its respondents are senior and experienced; the value of this evidence therefore lies in the depth and quality of practitioner judgement rather than in volume. Findings are accordingly reported qualitatively, and the study continues to claim analytical rather than statistical generalisation.')
P('Analytical techniques applied were: open coding of the interview text followed by grouping into themes aligned to the objectives; construct-level reading of the questionnaire responses to test whether the interview themes hold; content analysis of regulatory texts and public role profiles, coding activities as automated, augmented or unaffected and skills as technical, analytical, regulatory or behavioural; and triangulation of all streams before conclusions are drawn.')

# ---------------- 2
d.add_heading('2 Analysis of Primary Data',1)
d.add_heading('2.1 Profile of Respondents',2)
P('Interview respondents hold enterprise architecture, data architecture, technology and senior consulting roles supporting GRC service lines; questionnaire respondents work in internal audit, risk and GRC technology at analyst, manager and senior manager levels, all within banking service providers or global capability centres. Their work covers the UK, the EU, or both jurisdictions, and GRC-specific experience is split between early-career and mid-career practitioners. Reported AI maturity in their areas ranges from no AI use (with established non-AI automation) through pilots only, to use in some processes; no respondent reported extensive use. This spread is itself a finding: it suggests that, from the vantage point of practitioners and service providers, AI in banking GRC is still in a transitional phase rather than fully embedded.')

d.add_heading('2.2 Thematic Analysis of Interview Responses',2)
P('The following themes emerged from coding. They are presented under the secondary objective to which each primarily relates.')

d.add_heading('Objective (i): Activities automated or augmented and the changing demand for roles',3)
P('Theme 1 - Automation precedes AI, and the pilot-to-production gap is wide.',b=True)
P('Respondents consistently distinguished conventional automation from AI. Rules-based automation of regulatory monitoring and reporting, workflow automation in governance activities, control testing support and the aggregation of risk and compliance data for management reporting were described as established production capabilities. AI, by contrast, was described as largely confined to proofs of concept, pilots and targeted initiatives - intelligent document analysis, regulatory change assessment, enhanced risk scoring and compliance monitoring. A senior architect stated plainly:')
Q('I don\u2019t actually see it working in my experience. I have only ever seen pilot.')
P('Where AI was reported in production, the uses were specific and bounded: reconciliation and documentation of complex process flows, and identifying vulnerabilities and developing test scenarios to support operational resilience.')
P('Theme 2 - Routine, rules-based and data-intensive work is absorbed; judgement-intensive work resists.',b=True)
P('The activities most affected were data collection and consolidation for compliance reporting, control monitoring, regulatory reporting workflows, management information and dashboards, data mapping and design diagrams, and cross-format file comparison and reconciliation. The activities that resisted automation were challenging risk decisions, designing controls, agreeing remediation, providing assurance to senior stakeholders and regulators, and applying operations-specific knowledge of complex flows and regulation. A respondent captured the dividing line:')
Q('Tasks with structured inputs, defined rules and repeatable processes were generally easier to automate. Activities involving ambiguity, interpretation of regulatory intent, business context, or balancing competing risks continued to rely heavily on experienced practitioners.')
P('This is a direct, field-level confirmation of the substitution-complementarity distinction in Autor (2015) and of the routine-task exposure argument in Frey and Osborne (2017).')
P('Theme 3 - Roles are being recomposed, and some testing roles are disappearing.',b=True)
P('Respondents reported material time savings in data analysis, source-file comparison, cross-format reconciliation, business requirements interviewing, design, architecture diagramming and general development. A respondent observed that GRC testing roles had disappeared in their area. Together these responses indicate that the first roles to contract are those built around repetitive execution and testing, while the freed capacity is being redirected toward review, interpretation and higher-order design work - recomposition rather than simple elimination, again consistent with Autor (2015).')

d.add_heading('Objective (ii): Regulation, human oversight and competence',3)
P('Theme 4 - Human expert review remains non-negotiable.',b=True)
P('Every respondent who addressed oversight held that AI output requires expert human review before it can be used. Reasons given were that outputs are not reliably accurate word for word, that output quality depends on the quality of the input data, and that hallucination is a real risk. A respondent linked this directly to consequence:')
Q('Inaccurate insights and outputs can result in reputational damage and financial penalties.')
P('Meaningful challenge was described as requiring strong business and regulatory knowledge, an understanding of the underlying data, awareness of how AI can fail, and the confidence to apply judgement rather than accept the machine\u2019s conclusion.')
P('Theme 5 - An accountability gap at the top.',b=True)
P('The sharpest insight on oversight came from a senior enterprise architect working at the highest level of business operation:')
Q('C-level execs really do not understand what the AI is, and thus are likely to simply trust it.')
P('Others were less certain that accountable individuals are consistently equipped to supervise AI-assisted outputs. This matters because the SM&CR, DORA and SS1/23 all place accountability on named senior individuals or the management body. If those individuals cannot interrogate AI output, the regulatory model of human oversight is formally satisfied but substantively hollow - the automation-bias risk that the trust literature anticipates (Glikson and Woolley, 2020).')
P('Theme 6 - Regulation is felt unevenly in day-to-day practice.',b=True)
P('Views diverged. Some respondents had seen no practical regulatory effect on how AI is used in their area so far; another judged the regulatory environment to be inadequate in its ability to protect; others implied that regulatory exposure is what keeps human validation in place. The divergence suggests that regulatory obligations are being experienced mainly at the level of senior governance and model risk functions, and have not yet been translated into role-level operating practice for all practitioners - a gap the framework in Section 6 seeks to close.')

d.add_heading('Objective (iii): Skill gaps and reskilling strategies',3)
P('Theme 7 - A dual skill profile: technical fluency plus deep domain judgement.',b=True)
P('Skills reported as increasingly critical were AI, data, cyber, automation and resilience skills, with the largest gaps in AI risk, analytics and risk quantification. Alongside these, respondents emphasised reviewer skills - attention to detail and the ability to parse and review large volumes of information and generated content. The emerging GRC profile is therefore hybrid: enough technical literacy to understand what the tool is doing, and enough domain depth to know when it is wrong.')
P('Theme 8 - Loss of institutional knowledge is the most serious hidden risk.',b=True)
P('The most emphatic and repeated point in the primary data was that critical knowledge sits in the experience of long-serving subject matter experts, and that reducing headcount before that knowledge is captured undermines the very capability needed to validate AI. The respondent concerned returned to this point in several separate answers:')
Q('Too much business knowledge exists in people\u2019s heads, and if that knowledge is lost, it becomes much harder to validate, challenge and govern AI-generated outputs effectively. I would prioritise knowledge capture and transfer before any significant workforce reductions.')
P('This is an important contribution to the study. The literature reviewed treats automation as a question of which tasks are exposed; this finding shows that in GRC, the people whose routine work is most automatable may also hold the tacit knowledge that makes oversight of the automation possible. Sequencing - capture first, restructure second - therefore becomes a core workforce management decision.')
P('Theme 9 - Reskilling is a blend of build, buy and borrow, with specialist gaps persisting.',b=True)
P('Approaches reported were training existing staff, targeted external hiring, and reliance on technology and vendor teams. Training of existing staff was reported as effective for general tool use, but specialist gaps remain. Practical improvements suggested by respondents were dedicated learning modules for each regulation within each domain, and pre-embedded prompts within AI tools to shorten the path to productive use.')

d.add_heading('Objective (iv): Employee behavioural responses',3)
P('Theme 10 - Mixed responses: ease of use is welcomed, but anxiety and weak change management persist.',b=True)
P('Responses described growing confidence alongside concerns about job impact and output accuracy, which together sustain continued human validation. Where tools were accessible and easy to use, people embraced them. However, a senior respondent reported that many junior resources had become very worried in recent years, and judged business change management to be weak. Trust was described as contingent on the user\u2019s ability to understand and validate output; hallucination, lack of transparency and personal accountability all reduced willingness to rely on AI. Where management had done well, it had communicated the effectiveness of AI tools and demonstrated how they help in daily work.')
P('These responses map closely onto the reviewed constructs: perceived usefulness and ease of use driving adoption (Davis, 1989); transparency and reliability conditioning cognitive trust (Glikson and Woolley, 2020); and replaceability concern concentrated among junior staff, with implications for commitment and retention (Brougham and Haar, 2018).')

d.add_heading('2.3 Questionnaire: Validation of Interview Themes',2)
P('The questionnaire was used, as planned, to test whether interview themes hold across practitioners in risk, internal audit and GRC technology roles. Responses are read at construct level and reported directionally.')
T(['Construct','Direction of responses','Link to interview themes'],[
 ['A. Perceived usefulness','Generally positive: AI tools seen as improving performance and output quality; less consensus that they increase volume of work, and weakest where AI was only piloted.','Supports Themes 1 and 2 - value is real where AI is in production; limited where it is not.'],
 ['B. Perceived ease of use','Consistently positive on ease of use and learning effort; understanding what the tool is doing and how it reaches its outputs rated lower.','A usability-transparency gap: tools are easy to operate but not easy to see into (Themes 4 and 10).'],
 ['C. Trust in AI output','Reliability and the ability to defend an AI-assisted decision rated reasonably well; comfort relying on output without re-performing work and the effectiveness of override mechanisms were more guarded.','Confirms Theme 4 - practitioners keep a human check in place even when they rate outputs as reliable.'],
 ['D. Perceived replaceability','Strong agreement that roles will change substantially; concern about one\u2019s own role being performed by AI and about long-term career prospects was mixed.','Change is accepted as inevitable; anxiety is uneven, consistent with Theme 10 (concentrated among junior staff).'],
 ['E. Skills and reskilling','Data and analytical skills, and judgement and regulatory reasoning, both seen as more important; adequacy of organisational training was the weakest-rated item for some respondents.','Supports Themes 7 and 9 - the hybrid profile is recognised, but training provision lags.'],
 ['F. Oversight and accountability','Moderately positive on the knowledge of accountable persons; clarity of how regulatory oversight requirements are reflected in practice was more neutral.','Partially supports Themes 5 and 6 - oversight is present but not yet clearly operationalised.'],
 ['G. Organisational response','Communication and positive feeling about AI moderately positive; organisational commitment held up.','Indicates that clear communication can sustain engagement through change.']],
 widths=[3.5,6.5,6],cap='Table 2.1: Construct-level reading of questionnaire responses')
P('Open-ended responses reinforced the interview evidence: human review and oversight was identified as the part of GRC work that should not be automated, and dedicated modules for each regulation within every domain was identified as the single change that would most help staff adapt.')
P('Data quality note. A response returned the highest rating on every item from the trust section onward, including items that point in opposite directions (for example, strong agreement both with considering leaving the role because of AI and with feeling committed and positive). This pattern is characteristic of uniform responding and the response has been interpreted with caution; reverse-coded items will be used to screen for this in the final analysis, and reliability will be assessed using Cronbach\u2019s alpha as planned.',i=True)

# ---------------- 3
d.add_heading('3 Analysis of Secondary Data',1)
d.add_heading('3.1 Regulatory Requirements Bearing on Workforce Competence and Oversight (Objective ii)',2)
P('The regulatory texts were content-analysed to identify where they convert workforce competence and human oversight into explicit obligations. The analysis shows that, across both jurisdictions, regulation does not prescribe headcount but does prescribe who must understand what, who is accountable, and that capability must be evidenced.')
T(['Instrument','Workforce-relevant requirement','Implication for GRC workforce management'],[
 ['EU DORA','Management body bears ultimate responsibility for ICT risk and must keep knowledge and skills current through regular training; budget must cover ICT security awareness, resilience training and ICT skills for all staff; ICT risk must sit with an independent control function; ICT auditors must have sufficient ICT risk expertise; ICT security and resilience training is compulsory for all employees and senior management.','Board and senior-management capability becomes auditable. Training records, role-specific curricula and evidence of ICT expertise in audit and control functions become compliance artefacts, not HR extras.'],
 ['EU AI Act (as amended by the Digital Omnibus on AI, 2026)','AI literacy duty on all providers and deployers retained but recast as a duty of means - firms must take measures to support staff AI literacy, calibrated to role, context and the persons affected, without guaranteeing an individual level. Human oversight of high-risk systems must be effective: overseers must understand capabilities and limits, guard against automation bias, interpret outputs and be able to override. High-risk obligations (including credit scoring) were deferred by the Omnibus.','Role-based AI literacy programmes with documented evidence; designated, trained human overseers for high-risk use cases; the deferral creates a window to build oversight capability before obligations bite.'],
 ['PRA SS1/23','A relevant Senior Management Function holder must be accountable for the overall model risk framework, including ensuring effective challenge and independent validation; roles across the model lifecycle must be documented with the skills, experience and seniority required; validators need technical expertise, business familiarity and the standing to escalate. Applies to AI and machine-learning models.','Model validation and model risk roles must be resourced with people who understand AI techniques and the business; skills requirements must be written into role definitions; validators need organisational standing, not just technical skill.'],
 ['UK SM&CR (with FCA/PRA AI Update)','Regulators have chosen not to create a dedicated AI Senior Manager, relying instead on existing accountability: responsibility for technology systems normally sits with an existing SMF, and all Senior Managers are bound by the conduct rules, including taking reasonable steps.','Accountability for AI-assisted decisions falls on existing named executives - which makes the senior-level understanding gap identified in Theme 5 a personal regulatory exposure.'],
 ['FCA Consumer Duty','Firms must evidence good customer outcomes, including where decisions are automated.','Compliance and conduct-risk staff need the analytical skill to monitor outcome data produced by AI systems.'],
 ['EU AML package (AMLR / AMLA)','Requires a compliance manager at management-body level and a separate AML compliance officer, with defined duties and resourcing responsibility; a single directly applicable rulebook with direct EU supervision of selected entities.','Financial crime leadership roles are being redefined and split; senior AML roles must be able to own AI-enabled monitoring and screening outcomes.'],
 ['ECB supervisory priorities (2026-28)','Operational resilience and ICT capabilities is a standing priority, with a medium-term focus on banks\u2019 digital and AI-related strategies, governance and risk management, extending supervisory attention beyond prudential models to generative AI across the bank.','Supervisors will test whether AI governance is matched by people and skills - making workforce capability a supervisory, not only internal, concern.']],
 widths=[3,7,6],cap='Table 3.1: Mapping of UK and EU requirements to workforce competence and human oversight')
P('Supervisory survey evidence from the UK reinforces this picture. The joint Bank of England and FCA survey of AI in UK financial services reports that insufficient talent and access to skills is among the leading non-regulatory constraints on AI adoption; that many firms report only partial understanding of the AI technologies they use, largely because of reliance on third-party models; that risk and compliance is among the business areas with the highest reliance on third-party AI; and that accountability for AI is frequently split across several persons or bodies. Each of these findings converges with the primary evidence: specialist skill gaps (Theme 7), reliance on vendor teams (Theme 9), limited transparency (Theme 10) and diffuse senior accountability (Theme 5).')
P('An important interpretive point is that regulation acts simultaneously as a driver and a constraint. This extends the argument of Buckley et al. (2020): EU reporting and data rules drove RegTech adoption, and the newer oversight and literacy duties now drive investment in people. Where AI automates a task, regulation requires a capable human to remain answerable for it - which is why the First Review Report\u2019s proposition that a bank which automates without capable oversight has relocated rather than reduced its risk is supported by the texts themselves.')

d.add_heading('3.2 How GRC Job Descriptions Are Changing (Objectives i and iii)',2)
P('Publicly advertised GRC role profiles, recruiter market commentary and professional-body publications for UK and EU banking were reviewed and coded. The overall direction is clear: job descriptions are shifting from doing the work to supervising, challenging and improving the system that does the work. Table 3.2 summarises the shift for core GRC roles.')
T(['Role family','Traditional emphasis','Emerging emphasis in current role profiles'],[
 ['KYC / CDD analyst','Document collection, data entry, periodic file refresh, checklist completion.','Exception handling, reviewing AI-prepared case summaries, judging risk where data conflicts, quality-assuring automated refresh, escalation and customer outreach where required.'],
 ['Transaction monitoring / sanctions screening analyst','High-volume manual alert disposition, false-positive clearance.','Investigation of AI-prioritised alerts, typology knowledge, scenario and threshold tuning input, documenting rationale for overrides, feedback to models.'],
 ['Compliance testing / control assurance','Sample-based manual testing, evidence chasing, test scripting.','Continuous control monitoring design, full-population analytics, validating automated tests, interpreting control failures and root causes; pure execution testing roles contracting.'],
 ['Regulatory reporting','Data gathering, reconciliation, template population.','Data lineage and quality ownership, reviewing automated reconciliations, explaining variances to regulators.'],
 ['Compliance advisory / regulatory change','Reading and summarising rule changes, manual impact assessment.','Validating AI-generated regulatory horizon scanning and obligation mapping, interpreting regulatory intent, advising the business.'],
 ['Model risk / model validation','Statistical validation of credit and market risk models.','Validation of machine-learning and generative AI models, explainability, bias and fairness testing, LLM evaluation, third-party model assurance, ongoing performance and drift monitoring.'],
 ['Operational and ICT risk','Risk and control self-assessment facilitation, incident logs.','DORA-aligned ICT and third-party risk, resilience testing, AI incident management, concentration risk in AI and cloud providers.'],
 ['Internal audit','Cyclical assurance, manual sampling.','Data-led and continuous auditing, auditing AI systems and AI governance, greater advisory role, higher technology skill expectations at every level.']],
 widths=[3.5,5,7.5],cap='Table 3.2: Change in GRC role content as reflected in public role profiles')
P('New and rapidly growing roles. Alongside recomposed roles, a set of new role families now appears repeatedly in public postings and recruiter analyses for financial services:',al=WD_ALIGN_PARAGRAPH.LEFT)
B([('AI Governance Lead / Head of AI Risk: ','owns the responsible-AI framework, AI inventory, control design and ownership mapping, regulatory readiness and board reporting, coordinating risk, compliance, legal, audit and technology.'),
   ('AI Compliance Officer / Responsible AI Lead: ','translates the EU AI Act and supervisory expectations into policy, runs bias, fairness and explainability reviews, and approves use cases.'),
   ('AI Model Risk Validator: ','independent validation of machine-learning and generative AI models under model risk principles such as SS1/23.'),
   ('AI Assurance / AI Audit specialist: ','audits AI systems, their data, controls and oversight arrangements.'),
   ('Financial crime data scientist / FCC analytics lead: ','builds and tunes detection models and measures effectiveness and customer friction.'),
   ('Controls automation engineer / RegTech product owner: ','designs and maintains automated controls and compliance tooling, bridging GRC and engineering.'),
   ('ICT third-party and resilience risk manager: ','manages DORA obligations, including oversight of critical AI and cloud providers.'),
   ('Human-in-the-loop supervisor of AI agents: ','emerging in KYC and alert handling, supervising agentic workflows, handling exceptions and taking final decisions.')])
P('Skills most frequently requested. Content analysis of public role profiles for AI risk and compliance roles shows a consistent skills signal, coded here into the categories set out in the methodology:')
T(['Category','Skills appearing repeatedly in role profiles'],[
 ['Regulatory','EU AI Act, DORA, SS1/23 and model risk principles, SM&CR, Consumer Duty, AML rulebooks, data protection; frameworks such as the NIST AI Risk Management Framework and ISO/IEC 42001.'],
 ['Technical','Model risk management, AI and LLM evaluation, explainability tooling, bias and fairness testing, data lineage, privacy-by-design, vendor and third-party AI risk, AI incident response, red-teaming; working knowledge of SQL or Python and data visualisation for analytical roles.'],
 ['Analytical','Risk quantification, interpreting model outputs and performance metrics, root-cause analysis of control failures, full-population data analysis.'],
 ['Behavioural','Professional scepticism and effective challenge, translating technical issues for boards and regulators, stakeholder influence across lines of defence, change leadership, continuous learning.']],
 widths=[3,13],cap='Table 3.3: Skill requirements coded from public GRC role profiles')
P('Credentials such as AI governance, IT risk and IT audit certifications appear frequently but are generally listed as advantageous rather than mandatory. Recruiter commentary also notes that hybrid profiles combining regulatory knowledge with model-lifecycle fluency command a premium and are hard to source, and that leaders recruited from less regulated sectors often struggle with validation, documentation and audit expectations - reinforcing the case for building talent internally.')

d.add_heading('3.3 Workforce Management Practices Emerging in Banking GRC (Objectives iii and iv)',2)
B([('From pyramid to diamond. ','Industry commentary describes GRC teams moving from a broad base of junior processors toward a narrower base and a wider middle of experienced reviewers, investigators and specialists. This aligns with primary evidence that testing and processing roles contract first while review and design work grows - and raises the question of how future experts will be developed if entry-level roles shrink.'),
   ('Human effort redirected, not removed. ','Technology-provider and consulting analyses of agentic AI in KYC and AML consistently describe human effort shifting to exception handling, validation oversight, customer interaction and final decision-making, under a deliberate human-in-the-loop design.'),
   ('Build, buy and borrow. ','Banks are combining internal reskilling, targeted hiring of scarce AI-risk and data specialists, and reliance on vendors and capability centres. Professional research indicates that organisations which add new governance roles and actively rethink existing roles report greater benefit from AI than those that do not.'),
   ('Role-based AI literacy. ','Following the AI Act literacy duty and DORA training obligations, firms are moving from generic awareness training to tiered curricula: general users, operators, human overseers, model owners, and the management body - each with evidence of completion.'),
   ('Internal audit repositioning. ','The internal audit profession\u2019s forward vision describes auditors evolving into technology-enabled strategic advisers, with experienced and entry-level auditors alike requiring stronger technology skills.'),
   ('Skills-based and internal mobility models. ','Workforce planning is increasingly expressed in skills rather than job titles, with internal talent marketplaces used to move processing staff into investigation, quality assurance and data roles.'),
   ('Knowledge management as a control. ','Capturing procedures, regulatory interpretations and tacit expertise into structured knowledge bases is increasingly treated as a prerequisite for deploying AI - echoing Theme 8 directly.')])

# ---------------- 4
d.add_heading('4 Triangulation of Evidence',1)
P('Table 4.1 tests the convergence of the three evidence streams before conclusions are drawn.')
T(['Proposition','Primary evidence','Secondary evidence','Literature','Convergence'],[
 ['Routine, rules-based GRC tasks are automated first; judgement work resists.','Themes 2, 3','Role-profile shift from execution to review','Autor (2015); Frey & Osborne (2017)','Strong'],
 ['AI in GRC is largely augmentative and still maturing.','Theme 1; questionnaire maturity profile','Supervisory survey: widespread use but partial understanding','Arner et al. (2017)','Strong'],
 ['Regulation makes competence and oversight a compliance condition.','Themes 4, 6','DORA, AI Act, SS1/23, SM&CR, AMLR','Buckley et al. (2020)','Strong in texts; uneven in practice'],
 ['Senior accountable individuals may lack the understanding to challenge AI.','Theme 5; Construct F neutral','Split accountability; no dedicated AI SMF','Glikson & Woolley (2020)','Moderate to strong'],
 ['A hybrid technical-plus-domain skill profile is now required.','Theme 7; Construct E','Skills coded in Table 3.3','Cascio & Montealegre (2016)','Strong'],
 ['Institutional knowledge loss is a critical risk of AI-led restructuring.','Theme 8','Knowledge management as a deployment prerequisite','Extends Autor (2015)','Moderate - novel contribution'],
 ['Ease of use is high but transparency is low.','Theme 10; Constructs B and C','Explainability listed as a constraint and a skill','Davis (1989); Glikson & Woolley (2020)','Strong'],
 ['Replaceability concern is uneven and concentrated among junior staff; change management is weak.','Theme 10; Construct D','Pyramid-to-diamond shift','Brougham & Haar (2018)','Moderate']],
 widths=[4.2,3,3.3,3,2.5],cap='Table 4.1: Triangulation matrix')

# ---------------- 5
d.add_heading('5 Preliminary GRC Role-Impact Map',1)
T(['Category','GRC activities and roles'],[
 ['Automated (substitution)','Data collection and consolidation for reporting; cross-format reconciliation; management information and dashboard production; rules-based control monitoring; first-level alert triage and false-positive clearance; document classification and periodic KYC data refresh; execution-focused control testing.'],
 ['Augmented (complementarity)','Alert investigation; KYC and CDD risk judgement; regulatory change impact assessment; control design; model validation; audit planning and fieldwork; risk assessment; regulatory reporting review; compliance advisory - AI prepares, a human decides and is accountable.'],
 ['Newly created','AI governance lead; AI compliance officer / responsible-AI lead; AI model risk validator; AI assurance and audit specialist; financial crime data scientist; controls automation engineer / RegTech product owner; ICT third-party and resilience risk manager; human-in-the-loop supervisor of AI agents; knowledge and AI-literacy lead for GRC.'],
 ['Resistant (human-retained)','Interpreting regulatory intent; challenging risk decisions; agreeing remediation; assurance to boards and regulators; final accountability under SM&CR and DORA; ethical and conduct judgements.']],
 widths=[3.5,12.5],cap='Table 5.1: GRC role-impact map (preliminary)')

# ---------------- 6
d.add_heading('6 Preliminary Workforce Transformation Framework',1)
P('Integrating the validated themes, the regulatory mapping and the reviewed literature, a preliminary framework is proposed for GRC and HR leadership. It is sequenced deliberately, because the primary evidence shows that the order of decisions - particularly protecting knowledge before restructuring - determines whether oversight capability survives the transition.')
T(['Stage','Management decision','Key actions','Grounding'],[
 ['1. Diagnose','Which activities are exposed and which roles carry critical knowledge?','Activity-level mapping (automated / augmented / retained); AI inventory; identify knowledge-critical individuals.','Autor; Frey & Osborne; Themes 2, 8'],
 ['2. Protect knowledge','How is tacit expertise preserved before change?','Structured knowledge capture, procedure codification, mentoring and shadowing; no material reductions until capture is complete.','Theme 8; knowledge-management practice'],
 ['3. Redesign roles','Which roles are redesigned, reduced or created?','Rewrite role profiles toward review, challenge and system ownership; create AI governance, validation and assurance roles; protect a pipeline of developmental entry roles.','Tables 3.2, 5.1; Cascio & Montealegre'],
 ['4. Build oversight competence','Can accountable people meaningfully challenge AI?','Role-based AI literacy tiers including the board; competence matrix linked to DORA, AI Act, SS1/23 and SM&CR; designated trained overseers; evidence of training.','Table 3.1; Themes 4, 5'],
 ['5. Reskill and source','Build, buy or borrow?','Build hybrid skills internally through regulation-by-domain modules and embedded prompts; buy scarce AI-risk specialists; borrow vendor capacity under DORA third-party controls.','Themes 7, 9; Table 3.3'],
 ['6. Secure trust and adoption','Will staff rely on AI appropriately - neither blindly nor not at all?','Explainability at the point of use; clear override and challenge routes; demonstrate usefulness in daily work; measure appropriate reliance.','Davis; Glikson & Woolley; Theme 10'],
 ['7. Lead the change and retain talent','How are engagement and commitment sustained?','Transparent communication on role impact; visible career pathways from processing to investigation, QA and data roles; targeted support for junior staff; monitor commitment and turnover intent.','Brougham & Haar; Theme 10; Construct G']],
 widths=[2.6,3.6,6.4,3.4],cap='Table 6.1: Preliminary GRC workforce transformation framework')
P('Cross-cutting enablers are data quality and lineage (because output quality depends on input data), clear accountability mapping, and continuous monitoring of both AI performance and workforce indicators.')

# ---------------- 7
d.add_heading('7 Preliminary Implications for GRC and HR Leadership',1)
B(['Treat workforce capability as a regulatory control: competence and oversight evidence should be planned, budgeted and audited like any other control.',
   'Sequence restructuring behind knowledge capture; the cost of losing tacit expertise is borne later, in weaker oversight of AI.',
   'Close the senior-level understanding gap first; under SM&CR and DORA, accountability cannot be delegated to the tool or the vendor.',
   'Rewrite job descriptions around review, challenge and system ownership, and create the new governance and assurance roles explicitly rather than absorbing them informally.',
   'Invest in practical, regulation-specific and domain-specific learning rather than generic AI awareness.',
   'Strengthen business change management, with particular attention to junior staff whose roles are most exposed.',
   'Use the deferral of EU high-risk AI obligations as a capability-building window rather than a reason to delay.'])

d.add_heading('8 Scope Notes for the Next Stage',1)
P('The primary sample over-represents technology and architecture perspectives that work alongside GRC; additional responses from first-line compliance, financial crime operations and HR business partners would strengthen Objectives (iii) and (iv). The questionnaire reading is directional, and construct reliability will be confirmed in the final analysis. Regulatory findings are time-bound, given ongoing changes to the EU AI Act timetable and the UK accountability regime.')

# ---------------- refs
d.add_heading('Sources Consulted (Secondary Evidence)',1)
refs=['Bank of England and Financial Conduct Authority. (2024). Artificial intelligence in UK financial services - 2024.',
 'Bank of England, Prudential Regulation Authority. (2023). Supervisory Statement SS1/23: Model risk management principles for banks.',
 'Financial Conduct Authority. (2024). AI Update; and AI and the FCA: our approach.',
 'European Union. (2022). Regulation (EU) 2022/2554 on digital operational resilience for the financial sector (DORA).',
 'European Union. (2024). Regulation (EU) 2024/1689 (Artificial Intelligence Act).',
 'European Union. (2026). Regulation (EU) 2026/1744 (Digital Omnibus on AI), amending the Artificial Intelligence Act; European Commission, AI Omnibus enters into force.',
 'European Union. (2024). Regulation (EU) 2024/1624 (Anti-Money Laundering Regulation) and Regulation (EU) 2024/1620 establishing the Anti-Money Laundering Authority.',
 'European Central Bank, Banking Supervision. (2025). Supervisory priorities 2026-28.',
 'Thomson Reuters. (2025). Future of Professionals Report 2025.',
 'Internal Audit Foundation / The Institute of Internal Auditors. (2024). Internal Audit: Vision 2035 - Creating Our Future Together.',
 'Global Skill Development Council. (2026). AI Risk and Compliance Jobs Report 2026.',
 'IBM. (2026). Agentic AI is rewriting KYC and AML in banking.',
 'Christian & Timbers. (2026). Financial services AI executive search priorities 2026.',
 'Public GRC, risk, compliance and financial crime role profiles from UK and EU banking, reviewed on public job boards and AI-governance job boards (anonymised; no institution named).']
for i,r in enumerate(refs): d.add_paragraph(r,style='List Number')

d.save('/home/ubuntu/thesis/out/GRC_Workforce_Analysis_Narentheran_S.docx')
print('saved')
